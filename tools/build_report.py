from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Relatorio_Alinhamento_Sistema_RH.docx"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "20262E"
MUTED = "5B6573"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALL_OUT = "F4F6F9"
PALE_GREEN = "EAF4EE"
GREEN = "28603D"
PALE_GOLD = "FFF6DC"
GOLD = "7A5A00"
PALE_RED = "FCEBEC"
RED = "9B1C1C"
WHITE = "FFFFFF"
BORDER = "C9D1DA"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=TABLE_INDENT_DXA) -> None:
    assert sum(widths_dxa) == CONTENT_WIDTH_DXA, (widths_dxa, sum(widths_dxa))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_keep_with_next(paragraph, value=True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def set_keep_together(paragraph, value=True) -> None:
    paragraph.paragraph_format.keep_together = value


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_paragraph_border_bottom(paragraph, color=BLUE, size=10, space=5) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Title", 24, NAVY, 0, 4),
        ("Subtitle", 13, MUTED, 0, 16),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name not in ("Subtitle",)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_numbering(doc: Document) -> tuple[int, int, int]:
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_start = max(existing_abs or [0]) + 1
    num_start = max(existing_num or [0]) + 1

    def add_abstract(abstract_id: int, fmt: str, text: str) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        nsid = OxmlElement("w:nsid")
        nsid.set(qn("w:val"), f"{abstract_id:08X}")
        abstract.append(nsid)
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
        r_pr.append(fonts)
        lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

    add_abstract(abs_start, "bullet", "•")
    add_abstract(abs_start + 1, "decimal", "%1.")
    # Use two independent decimal numbering instances so each numbered section
    # can restart at 1 instead of Word continuing the earlier sequence.
    for offset, abs_id in enumerate((abs_start, abs_start + 1, abs_start + 1)):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_start + offset))
        abstract_num_id = OxmlElement("w:abstractNumId")
        abstract_num_id.set(qn("w:val"), str(abs_id))
        num.append(abstract_num_id)
        if offset == 2:
            level_override = OxmlElement("w:lvlOverride")
            level_override.set(qn("w:ilvl"), "0")
            start_override = OxmlElement("w:startOverride")
            start_override.set(qn("w:val"), "1")
            level_override.append(start_override)
            num.append(level_override)
        numbering.append(num)
    return num_start, num_start + 1, num_start + 2


def add_list_item(doc: Document, text: str, num_id: int, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_status_chip(cell, text: str, kind: str) -> None:
    colors = {
        "confirmado": (PALE_GREEN, GREEN),
        "pendente": (PALE_GOLD, GOLD),
        "risco": (PALE_RED, RED),
        "proposta": (BLUE_GRAY, DARK_BLUE),
    }
    fill, color = colors[kind]
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=9.5, color=color, bold=True)


def style_table_header(row) -> None:
    set_repeat_table_header(row)
    for cell in row.cells:
        set_cell_shading(cell, LIGHT_GRAY)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                set_run_font(run, size=9.5, color=NAVY, bold=True)


def style_table_body(table, font_size=9.5) -> None:
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    set_run_font(run, size=font_size, color=INK)


def add_table(doc, headers, rows, widths_dxa, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            row.cells[i].text = str(value)
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    style_table_header(table.rows[0])
    style_table_body(table, font_size=font_size)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, title: str, body: str, kind="info"):
    palette = {
        "info": (CALL_OUT, NAVY),
        "confirmed": (PALE_GREEN, GREEN),
        "pending": (PALE_GOLD, GOLD),
        "risk": (PALE_RED, RED),
    }
    fill, accent = palette[kind]
    table = doc.add_table(rows=1, cols=1)
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    cell.text = ""
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=accent, size=8)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_body(doc, text, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def setup_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("SISTEMA RH - ANÁLISE CURRICULAR")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r1 = fp.add_run("Alinhamento interno  |  Página ")
    set_run_font(r1, size=9, color=MUTED)
    add_page_number(fp)


def add_masthead(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("RELATÓRIO DE ALINHAMENTO")
    set_run_font(r, size=10, color=BLUE, bold=True)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("Sistema Web de RH para Recrutamento por Posto")
    set_run_font(r, size=24, color=NAVY, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    r = subtitle.add_run("Decisões confirmadas, situação atual, escopo proposto e pendências")
    set_run_font(r, size=13, color=MUTED)

    metadata = [
        ("Destinatários", "Equipe de RH, gestão e tecnologia"),
        ("Data de referência", "1º de setembro de 2026"),
        ("Status", "Documento para alinhamento interno - sistema ainda não implementado"),
        ("Prioridade", "Reduzir o tempo de análise de mais de 1.000 currículos por mês"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=10.5, color=NAVY, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color=INK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(10)
    set_paragraph_border_bottom(rule, color=BLUE, size=12, space=4)


def build() -> None:
    doc = Document()
    doc.core_properties.title = "Relatório de Alinhamento - Sistema Web de RH"
    doc.core_properties.subject = "Decisões confirmadas, escopo proposto e pendências"
    doc.core_properties.author = "Equipe do Projeto Sistema RH"
    doc.core_properties.keywords = "RH, recrutamento, currículos, Weboper, inteligência artificial"
    setup_styles(doc)
    setup_page(doc)
    bullet_id, number_id, next_steps_number_id = add_numbering(doc)
    add_masthead(doc)

    add_callout(
        doc,
        "Resumo executivo",
        "O projeto pretende centralizar currículos, extrair informações com IA e organizar a triagem por requisitos profissionais e deslocamento até o posto. O foco do primeiro MVP é reduzir o trabalho manual de duas pessoas do RH diante de mais de 1.000 currículos por mês. Até esta data, somente o banco do Weboper foi inspecionado; não existe aplicação, repositório ou infraestrutura implantada.",
        "info",
    )

    add_heading(doc, "1. Objetivo do projeto", 1)
    add_body(
        doc,
        "Construir um sistema web de recrutamento vinculado aos postos de trabalho atendidos pela empresa. O sistema deverá centralizar candidaturas, preservar os currículos originais, extrair dados relevantes, comparar candidatos com requisitos objetivos da vaga e apoiar a avaliação de deslocamento, mantendo a decisão final com o RH.",
    )
    add_body(
        doc,
        "O principal problema operacional confirmado é o tempo gasto para abrir e analisar currículos individualmente, inclusive apenas para descobrir onde o candidato mora. A localização pode aparecer como endereço completo, CEP, bairro, município, texto incompleto ou não aparecer.",
    )

    add_heading(doc, "2. Decisões e informações confirmadas", 1)
    confirmed_rows = [
        ("Volume", "Mais de 1.000 currículos por mês."),
        ("Usuários", "No máximo duas pessoas do RH utilizando o sistema simultaneamente."),
        ("Prioridade do MVP", "Reduzir o tempo de análise e triagem dos currículos."),
        ("Postos", "Usar somente registros ativos da tabela CAD_CLIENTE do Weboper."),
        ("Weboper", "Consulta somente de leitura; dados de RH permanecerão em banco próprio."),
        ("Decisão humana", "A IA apoiará o RH, mas não contratará nem rejeitará candidatos autonomamente."),
        ("Localização", "Qualificação profissional e deslocamento serão avaliados separadamente."),
        ("Dados ausentes", "Não informado não será tratado como não atende."),
        ("Rotas", "Rotas e custos de transporte entram após validar o fluxo principal em piloto."),
    ]
    table = add_table(doc, ["Tema", "Definição confirmada"], confirmed_rows, [2200, 7160], font_size=9.7)
    for row in table.rows[1:]:
        set_cell_shading(row.cells[0], PALE_GREEN)
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, size=9.7, color=GREEN, bold=True)

    add_heading(doc, "3. Situação técnica verificada", 1)
    add_body(
        doc,
        "Foi acessada a sessão já autenticada do phpMyAdmin para inspeção da estrutura e execução de consultas SELECT limitadas. Nenhuma inserção, alteração ou exclusão foi realizada.",
    )
    technical_rows = [
        ("Fonte dos postos", "CAD_CLIENTE"),
        ("Quantidade total", "300 registros"),
        ("Postos ativos", "134 registros"),
        ("Postos inativos", "165 registros"),
        ("Sem situação", "1 registro"),
        ("Chave", "CHAVE - int, chave primária e AUTO_INCREMENT"),
        ("Índice adicional", "NOME_FANTASIA"),
        ("Campos necessários", "CHAVE, RAZAO_SOCIAL, NOME_FANTASIA, SITUACAO, ENDERECO, BAIRRO, MUNICIPIO, UF e CEP"),
        ("Exemplo validado", "LEBLON POWER - identificador 252 - posto ativo no Leblon"),
    ]
    add_table(doc, ["Item", "Resultado verificado"], technical_rows, [2400, 6960], font_size=9.7)

    add_heading(doc, "3.1 Qualidade dos endereços", 2)
    quality_rows = [
        ("Sem endereço", "21"),
        ("Sem bairro", "22"),
        ("Sem município", "21"),
        ("Sem UF", "41"),
        ("Sem CEP", "43"),
    ]
    add_table(doc, ["Pendência cadastral", "Quantidade"], quality_rows, [6500, 2860], font_size=9.7)
    add_callout(
        doc,
        "Impacto no sistema",
        "Endereços incompletos não devem impedir a sincronização do posto, mas precisam aparecer como pendência antes do cálculo de rota. A aplicação também deverá validar a conversão dos textos atuais de latin1 para UTF-8.",
        "pending",
    )

    add_heading(doc, "3.2 Situação de segurança", 2)
    add_list_item(doc, "Uma credencial real do banco foi compartilhada durante o levantamento e deve ser substituída antes da implementação.", bullet_id)
    add_list_item(doc, "Ainda não foi criado um usuário MySQL exclusivo com permissão somente de leitura.", bullet_id)
    add_list_item(doc, "A conexão da futura aplicação com o MySQL não foi configurada nem testada.", bullet_id)
    add_list_item(doc, "O diretório do projeto está vazio: não há repositório, código, protótipo ou infraestrutura criada.", bullet_id)

    add_heading(doc, "4. Canais atuais de recebimento de currículos", 1)
    add_body(doc, "O RH confirmou cinco canais de entrada. Cada um exigirá um adaptador próprio, mas todos deverão terminar na mesma fila de importação e análise.")
    channel_rows = [
        ("RioVagas", "Currículos enviados para uma conta Gmail criada pelo RH.", "Importação por Gmail/IMAP ou API, após validar conta e pasta."),
        ("People", "Sistema de recrutamento que distribui vagas para Google for Jobs, LinkedIn, Netvagas e Indeed.", "Verificar API, webhook, exportação e identificação da origem."),
        ("Indeed", "Conta própria gratuita, com funções limitadas.", "Inspecionar recursos disponíveis; não presumir acesso à API."),
        ("WhatsApp", "Currículos avulsos, indicações de funcionários e público geral.", "Upload manual no MVP; API oficial apenas em evolução futura."),
        ("Currículo físico", "Documentos entregues na sede ou no RH.", "Digitalização/foto, OCR, preservação do original e revisão humana."),
    ]
    add_table(doc, ["Canal", "Funcionamento atual", "Tratamento recomendado"], channel_rows, [1500, 3570, 4290], font_size=9.1)

    add_heading(doc, "4.1 Regras comuns de importação", 2)
    common_rules = [
        "Registrar canal, mensagem ou identificação de origem, data de recebimento e vaga relacionada.",
        "Preservar o arquivo original e manter histórico de reprocessamento.",
        "Calcular hash do arquivo antes de acionar a IA para reduzir duplicidades e custos.",
        "Não unir pessoas somente pelo nome; combinar arquivo, identificador externo, e-mail, telefone e conteúdo normalizado.",
        "Enviar arquivos ilegíveis, incompletos ou com baixa confiança para revisão do RH.",
        "Não mover, apagar ou responder mensagens automaticamente na primeira versão.",
    ]
    for item in common_rules:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "5. Fluxo funcional recomendado para o MVP", 1)
    flow_steps = [
        "Receber a candidatura por Gmail, plataforma, upload manual ou digitalização.",
        "Registrar a origem e preservar o currículo original.",
        "Verificar arquivo repetido e possível candidato duplicado.",
        "Extrair texto, dados profissionais e localização, mantendo evidência e página de origem.",
        "Normalizar endereço, CEP, bairro ou município sem inventar informações ausentes.",
        "Associar a candidatura à vaga informada ou colocá-la na fila de vaga pendente.",
        "Comparar requisitos profissionais e deslocamento em blocos independentes.",
        "Apresentar a fila ordenada ao RH para revisão e próxima ação.",
    ]
    for step in flow_steps:
        add_list_item(doc, step, number_id)

    add_callout(
        doc,
        "Regra de decisão",
        "O sistema poderá ordenar a fila para que o RH comece por candidatos profissionalmente compatíveis e com deslocamento mais viável. Não haverá reprovação automática por bairro, município, região ou custo de transporte.",
        "confirmed",
    )

    add_heading(doc, "6. Informações exibidas na triagem", 1)
    triage_rows = [
        ("Qualificação profissional", "Requisitos atendidos, evidências, requisitos não informados e incompatibilidades a confirmar."),
        ("Localização", "Texto original extraído, endereço/CEP/bairro/município normalizado e nível de precisão."),
        ("Deslocamento", "Curto, moderado, longo ou desconhecido; classificação transparente e separada da qualificação."),
        ("Origem", "RioVagas, People, Indeed, WhatsApp, físico ou outro canal futuro."),
        ("Qualidade da extração", "Confirmada, pendente de revisão, ilegível ou incompleta."),
        ("Duplicidade", "Novo candidato, possível duplicidade ou arquivo já processado."),
        ("Processo seletivo", "Novo, triagem, contato, entrevista, proposta, contratado ou encerrado."),
    ]
    add_table(doc, ["Bloco", "Conteúdo"], triage_rows, [2400, 6960], font_size=9.5)

    add_heading(doc, "7. Telas propostas - pendentes de aprovação formal", 1)
    add_callout(
        doc,
        "Status desta seção",
        "O mapa abaixo foi recomendado durante a conversa, mas ainda não houve aprovação formal do conjunto completo de telas.",
        "pending",
    )
    screen_rows = [
        ("Login", "Acesso seguro, recuperação de senha e autenticação em dois fatores."),
        ("Hoje no RH", "Currículos novos, revisões pendentes, vagas urgentes, falhas e próximos contatos."),
        ("Vagas", "Lista, criação, edição, requisitos, escala, salário, benefícios e posto ativo."),
        ("Detalhes da vaga", "Resumo, candidatos, etapas, canais de publicação e histórico."),
        ("Triagem inteligente", "Fila principal com qualificação, localização, deslocamento, origem e pendências."),
        ("Perfil do candidato", "Currículo original ao lado dos dados extraídos, corrigidos e confirmados."),
        ("Banco de candidatos", "Pesquisa por experiência, cursos, localização, disponibilidade e histórico."),
        ("Importações", "Gmail/RioVagas, People, Indeed, WhatsApp, físicos, duplicidades e erros."),
        ("Postos", "Somente postos ativos sincronizados do Weboper e qualidade do endereço."),
        ("Administração", "Usuários, integrações, custos, auditoria, privacidade e configurações."),
        ("Área pública", "Lista de vagas, detalhes, formulário de candidatura e confirmação."),
    ]
    add_table(doc, ["Tela", "Finalidade"], screen_rows, [2200, 7160], font_size=9.5)

    add_heading(doc, "8. Escopo recomendado por fase", 1)
    phase_rows = [
        ("0. Segurança e acessos", "Trocar credencial exposta, criar usuário MySQL somente leitura, definir segredos e validar backup."),
        ("1. Base do MVP", "Autenticação, postos ativos, vagas, candidatos, arquivos originais e histórico."),
        ("2. Entradas prioritárias", "Upload manual, currículo físico e Gmail/RioVagas; testar deduplicação e reprocessamento."),
        ("3. IA e triagem", "Extração estruturada, localização, evidências, comparação com vaga e revisão humana."),
        ("4. Integrações externas", "People e Indeed conforme API, plano, contrato e recursos disponíveis."),
        ("5. Deslocamento", "Rotas de transporte público, horários, conexões, tarifas e custo por escala."),
        ("6. Evoluções", "Agenda, indicadores avançados, multipublicação e formulário público ampliado."),
    ]
    table = add_table(doc, ["Fase", "Entregas"], phase_rows, [2200, 7160], font_size=9.5)
    for row in table.rows[1:]:
        set_cell_shading(row.cells[0], BLUE_GRAY)
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, size=9.5, color=DARK_BLUE, bold=True)

    add_heading(doc, "9. Pendências para decisão ou levantamento", 1)
    pending_rows = [
        ("Alta", "Credencial do Weboper", "Substituir a credencial exposta e criar usuário somente leitura."),
        ("Alta", "People", "Identificar nome completo/URL, plano contratado, API, exportação e forma de receber candidaturas."),
        ("Alta", "Gmail do RioVagas", "Confirmar conta, pasta, volume, histórico a importar e formatos recebidos."),
        ("Alta", "Primeiro piloto", "Escolher função e vaga que representarão o teste inicial."),
        ("Alta", "Critérios da vaga", "Definir requisitos objetivos, obrigatórios e desejáveis antes da análise."),
        ("Média", "Indeed", "Inspecionar recursos da conta gratuita e possibilidades autorizadas de integração/exportação."),
        ("Média", "WhatsApp", "Confirmar se é conta pessoal ou WhatsApp Business e definir fluxo operacional do upload."),
        ("Média", "Currículos físicos", "Definir scanner/celular, responsável e procedimento de guarda ou descarte do papel."),
        ("Média", "KingHost", "Confirmar se a caixa KingHost ainda participa de algum fluxo de recrutamento."),
        ("Média", "Tecnologia", "Definir stack, hospedagem, banco próprio, armazenamento de arquivos e orçamento."),
        ("Média", "OpenAI", "Selecionar modelo após piloto, estimar custos e definir limites/reprocessamentos."),
        ("Posterior", "Mapas", "Escolher fornecedor e validar transporte público, cobertura, tarifas e custos."),
        ("Jurídica", "LGPD", "Definir finalidade, base legal, transparência, retenção, fornecedores e direitos dos candidatos."),
        ("Visual", "Identidade e UX", "Aprovar mapa de telas, identidade visual e detalhes da tela de triagem."),
    ]
    table = add_table(doc, ["Prioridade", "Tema", "O que falta"], pending_rows, [1200, 2400, 5760], font_size=9.1)
    for row in table.rows[1:]:
        label = row.cells[0].text.strip()
        kind = "risco" if label in ("Alta", "Jurídica") else "pendente"
        fill, color = (PALE_RED, RED) if kind == "risco" else (PALE_GOLD, GOLD)
        set_cell_shading(row.cells[0], fill)
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, size=9.1, color=color, bold=True)

    add_heading(doc, "10. Propostas ainda não aprovadas", 1)
    proposals = [
        "Sincronizar os postos ativos periodicamente e preservar no histórico os postos que depois se tornarem inativos.",
        "Usar formulário próprio como entrada principal em canais que aceitem link externo.",
        "Utilizar LinkedIn por link externo no MVP e avaliar Apply Connect somente após aprovação comercial.",
        "Fazer upload manual de currículos do WhatsApp na primeira versão.",
        "Executar comparação detalhada com IA somente quando o candidato estiver associado a uma vaga.",
        "Manter uma única fila de importação para todos os canais, com adaptadores independentes.",
        "Começar pelas telas de Vagas, Triagem, Perfil do candidato, Importações e formulário público.",
        "Avaliar sincronização completa dos 134 postos ativos em intervalos regulares, pois não há coluna de última atualização no Weboper.",
    ]
    for item in proposals:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "11. Riscos principais e controles esperados", 1)
    risk_rows = [
        ("Acesso indevido a currículos", "Permissões por função, autenticação forte, auditoria e acesso mínimo necessário."),
        ("Vazamento de credenciais", "Segredos somente no servidor, rotação de senha e nenhuma credencial em código ou conversa."),
        ("Duplicidade entre canais", "Hash, identificador externo, contato e revisão antes de unir cadastros."),
        ("Erro ou invenção da IA", "Saída estruturada, evidências por trecho/página, validação e revisão humana."),
        ("Viés na seleção", "Critérios profissionais objetivos; não usar foto, raça, religião, família, nome ou idade genérica."),
        ("Endereço impreciso", "Preservar texto original, indicar precisão e não reprovar automaticamente."),
        ("Arquivo malicioso ou excessivo", "Limite de tamanho, tipos permitidos, verificação de conteúdo e processamento isolado."),
        ("Custo de APIs", "Deduplicação antes da IA, limites, métricas, alertas e processamento sob demanda."),
        ("Falha de integração", "Fila durável, retentativas controladas, registro de erro e reprocessamento pelo RH."),
        ("Retenção inadequada", "Política jurídica de retenção, descarte seguro e atendimento aos direitos do candidato."),
    ]
    add_table(doc, ["Risco", "Controle esperado"], risk_rows, [3100, 6260], font_size=9.3)

    add_heading(doc, "12. Próximos passos recomendados", 1)
    next_steps = [
        "Trocar a credencial do MySQL que foi exposta e criar uma conta exclusiva somente leitura para CAD_CLIENTE.",
        "Identificar exatamente o sistema People e levantar suas opções de API, exportação e recebimento de candidatos.",
        "Visitar o RH para mapear Gmail, People, Indeed, WhatsApp e rotina dos currículos físicos sem compartilhar dados pessoais desnecessários.",
        "Escolher uma vaga/função piloto e formalizar seus requisitos objetivos, escala, horários e endereço do posto.",
        "Aprovar o mapa de telas e o fluxo principal de triagem antes de iniciar a implementação.",
        "Definir stack, hospedagem, orçamento e responsabilidades de segurança e LGPD.",
        "Escrever a especificação técnica validada e somente depois elaborar o plano de implementação.",
    ]
    for item in next_steps:
        add_list_item(doc, item, next_steps_number_id)

    add_callout(
        doc,
        "Estado do projeto em 01/09/2026",
        "Levantamento e desenho em andamento. O banco do Weboper foi inspecionado, mas nenhuma integração do sistema, tela, repositório, serviço de IA, importador de e-mail ou infraestrutura foi implementado ou testado.",
        "risk",
    )

    add_heading(doc, "Apêndice A - Referências preliminares de integração", 1)
    add_body(doc, "As referências abaixo sustentam apenas o levantamento de viabilidade. A contratação, elegibilidade e disponibilidade para a conta da empresa ainda precisam ser confirmadas.")
    refs = [
        ("LinkedIn - fluxo de candidatura e redirecionamento externo", "https://www.linkedin.com/help/linkedin/answer/a414058/application-flow-for-members-applying-for-your-job-posting?lang=en-us"),
        ("LinkedIn - Apply Connect", "https://learn.microsoft.com/en-us/linkedin/talent/apply-connect"),
        ("Indeed - Job Sync API e Indeed Apply", "https://docs.indeed.com/job-sync-api/integrate-with-job-sync-api"),
        ("RioVagas - publicação e recebimento por e-mail", "https://riovagas.com.br/vagas/"),
        ("Catho - portal de integrações", "https://desenvolvedores.catho.com.br/developers-integration-portal"),
        ("InfoJobs/Pandapé - integração por API", "https://supportcenter.infojobs.com.br/hc/pt-br/articles/41245648835601"),
        ("Vagas for Business - APIs", "https://forbusiness.vagas.com.br/developers"),
    ]
    for label, url in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run("• ")
        set_run_font(r, color=BLUE, bold=True)
        add_hyperlink(p, label, url)

    add_heading(doc, "Apêndice B - Legenda de status", 1)
    legend = doc.add_table(rows=5, cols=2)
    legend.rows[0].cells[0].text = "Status"
    legend.rows[0].cells[1].text = "Interpretação"
    legend.rows[1].cells[0].text = "CONFIRMADO"
    legend.rows[1].cells[1].text = "Informação fornecida pelo responsável/RH ou verificada diretamente no banco."
    legend.rows[2].cells[0].text = "PROPOSTA"
    legend.rows[2].cells[1].text = "Recomendação técnica ou funcional ainda sujeita à validação."
    legend.rows[3].cells[0].text = "PENDENTE"
    legend.rows[3].cells[1].text = "Informação, decisão, acesso ou teste ainda necessário."
    legend.rows[4].cells[0].text = "NÃO IMPLEMENTADO"
    legend.rows[4].cells[1].text = "Funcionalidade apenas discutida; ainda não existe no ambiente."
    set_table_geometry(legend, [2400, 6960])
    set_table_borders(legend)
    style_table_header(legend.rows[0])
    style_table_body(legend, font_size=9.7)
    for idx, kind in enumerate(("confirmado", "proposta", "pendente", "risco")):
        row_idx = idx + 1
        text = legend.rows[row_idx].cells[0].text
        legend.rows[row_idx].cells[0].text = ""
        add_status_chip(legend.rows[row_idx].cells[0], text, kind)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
