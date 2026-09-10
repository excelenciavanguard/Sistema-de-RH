import io

from docx import Document

from app.services.text_extraction import extract_text


def test_extracts_utf8_txt():
    text, note = extract_text("Experiência em limpeza\nEnsino médio completo".encode("utf-8"), "txt")
    assert "Experiência em limpeza" in text
    assert note is None


def test_extracts_docx_paragraphs():
    document = Document()
    document.add_paragraph("Rafael Santos")
    document.add_paragraph("Auxiliar de serviços gerais")
    buffer = io.BytesIO()
    document.save(buffer)

    text, note = extract_text(buffer.getvalue(), "docx")
    assert "Rafael Santos" in text
    assert "Auxiliar de serviços gerais" in text
    assert note is None


def test_doc_legacy_is_sent_to_review():
    text, note = extract_text(b"legacy", "doc")
    assert text is None
    assert "revisão" in note

